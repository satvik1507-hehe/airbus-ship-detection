"""Builds the documents in docs/ from the numbers in results/.

  python docs/build_docs.py

Produces
  Team9_Project35_Writeup.pdf     two-page summary
  Team9_Project35_Slides.pptx     presentation deck
  Team9_Project35_Viva_Guide.pdf  explanation of the project + likely viva questions
"""
import json
import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
RES, FIG, DOCS = ROOT / "results", ROOT / "results" / "figures", ROOT / "docs"

TEAM = "Team 9  |  Project No. 35"
MEMBERS = "Rohan Suresh (PES1UG24CS383)  ·  Satvik L (PES1UG24CS423)"
COURSE = "UE24CS352A – Machine Learning  |  Mini-Project"
TITLE = "Airbus Ship Detection: Traditional ML vs. Convolutional Neural Networks"

# ----------------------------------------------------------------- numbers
trad = json.load(open(RES / "traditional_results.json"))
cnn = json.load(open(RES / "cnn_results.json"))
table1 = pd.read_csv(RES / "table1_test_accuracy.csv", index_col=0)
hist = {k: json.load(open(RES / f"{k}_history.json")) for k in ("sim", "tl")}
labels = pd.read_csv(ROOT / "data" / "labels.csv")

N = len(labels)
N_TR, N_DV, N_TE = trad["n_train"], trad["n_dev"], trad["n_test"]
SHIP = labels.has_ship.mean()
BASE = trad["majority_baseline"]
ALGS = ["LDA", "KNN", "NB", "RF", "SVM"]
ALL = "w/ all"
RAW = "w/o (raw pixels)"
acc_all = {a: trad["detail"][f"{a}|{ALL}"]["accuracy"] for a in ALGS}
rf = trad["detail"][f"RF|{ALL}"]
best_trad = max(acc_all, key=acc_all.get)
sim, tl = cnn["sim"]["test_at_best_threshold"], cnn["tl"]["test_at_best_threshold"]
sim5, tl5 = cnn["sim"]["test_at_0.5"], cnn["tl"]["test_at_0.5"]
sim_auc, tl_auc = cnn["sim"]["test_roc_auc"], cnn["tl"]["test_roc_auc"]
cv = trad["cv"]
best_overall = max({**acc_all, "TL-CNN": tl["accuracy"], "SIM-CNN": sim["accuracy"]}.items(),
                   key=lambda kv: kv[1])
best_cell = table1.stack().idxmax()
P_SIM, P_TL = hist["sim"]["trainable_parameters"], hist["tl"]["trainable_parameters"]
below_base = [a for a in ALGS if acc_all[a] < BASE]


def pct(x, d=1):
    return f"{100 * x:.{d}f}%"


cnn_order = ("SIM-CNN", "TL-CNN") if sim["accuracy"] >= tl["accuracy"] else ("TL-CNN", "SIM-CNN")
cnn_hi, cnn_lo = max(sim["accuracy"], tl["accuracy"]), min(sim["accuracy"], tl["accuracy"])
CNN_SENTENCE = (f"{cnn_order[0]} ({pct(cnn_hi)}) scored higher than {cnn_order[1]} ({pct(cnn_lo)}) "
                f"on our test set")
FE_GAIN = {a: table1.loc[a, ALL] - table1.loc[a, RAW] for a in ALGS}
fe_helped = [a for a in ALGS if FE_GAIN[a] > 0.01]
fe_hurt = [a for a in ALGS if FE_GAIN[a] < -0.01]

# ----------------------------------------------------------------- pdf helpers
ss = getSampleStyleSheet()


def styles(base=9.6):
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontSize=base + 5.5, leading=base + 8,
                                spaceAfter=2),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontSize=base - 0.6, leading=base + 2,
                              alignment=TA_CENTER, textColor=colors.HexColor("#444444")),
        "h": ParagraphStyle("h", parent=ss["Heading2"], fontSize=base + 1.6, leading=base + 4,
                            spaceBefore=6, spaceAfter=2, textColor=colors.HexColor("#0b3d6b")),
        "h1": ParagraphStyle("h1", parent=ss["Heading1"], fontSize=base + 5, leading=base + 8,
                             spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#0b3d6b")),
        "b": ParagraphStyle("b", parent=ss["Normal"], fontSize=base, leading=base + 3.2,
                            spaceAfter=3),
        "q": ParagraphStyle("q", parent=ss["Normal"], fontSize=base, leading=base + 3.2,
                            spaceBefore=5, spaceAfter=1, fontName="Helvetica-Bold"),
        "cap": ParagraphStyle("c", parent=ss["Normal"], fontSize=base - 1.4, leading=base + 1,
                              alignment=TA_CENTER, textColor=colors.HexColor("#444444")),
        "code": ParagraphStyle("code", parent=ss["Code"], fontSize=base - 1, leading=base + 2,
                               backColor=colors.HexColor("#f2f2f2"), borderPadding=4,
                               spaceBefore=3, spaceAfter=6),
    }


def bullets(items, st, left=12):
    return ListFlowable([ListItem(Paragraph(i, st), leftIndent=left) for i in items],
                        bulletType="bullet", start="•", leftIndent=left, bulletFontSize=7)


def grid(data, col_widths=None, font=8.2, header=True, bold_cells=()):
    t = Table(data, colWidths=col_widths, hAlign="CENTER")
    st = [("FONTSIZE", (0, 0), (-1, -1), font), ("LEADING", (0, 0), (-1, -1), font + 2.2),
          ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
          ("ALIGN", (1, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("TOPPADDING", (0, 0), (-1, -1), 1.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)]
    if header:
        st += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbe7f3")),
               ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
    for (r, c) in bold_cells:
        st.append(("FONTNAME", (c, r), (c, r), "Helvetica-Bold"))
    t.setStyle(TableStyle(st))
    return t


def table1_rows():
    cols = list(table1.columns)
    head = ["", "His", "Ha", "Hu", "His+Ha", "Ha+Hu", "Hu+His", "All three", "Raw pixels"]
    rows = [head] + [[a] + [f"{table1.loc[a, c]:.3f}" for c in cols] for a in ALGS]
    r, c = ALGS.index(best_cell[0]) + 1, cols.index(best_cell[1]) + 1
    return rows, [(r, c)]


def summary_rows():
    rows = [["Model", "Accuracy", "Precision", "Recall", "Specificity / notes"]]
    for a in ALGS:
        d = trad["detail"][f"{a}|{ALL}"]
        rows.append([f"{a} (all features)", pct(d["accuracy"]), pct(d["precision"]),
                     pct(d["recall"]), f"10-fold CV {pct(cv[a]['mean'])} ± {pct(cv[a]['std'])}"])
    rows.append([f"TL-CNN (thr {tl['threshold']})", pct(tl["accuracy"]), pct(tl["precision"]),
                 pct(tl["recall"]), f"specificity {pct(tl['specificity'])}, AUC {tl_auc:.3f}"])
    rows.append([f"SIM-CNN (thr {sim['threshold']})", pct(sim["accuracy"]), pct(sim["precision"]),
                 pct(sim["recall"]), f"specificity {pct(sim['specificity'])}, AUC {sim_auc:.3f}"])
    rows.append(["Always predict 'Ship'", pct(BASE), pct(SHIP), "100.0%", "specificity 0% (baseline)"])
    return rows


def page_number(canvas, doc):
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawCentredString(A4[0] / 2, 0.8 * cm, f"{TEAM}  –  page {doc.page}")


# ----------------------------------------------------------------- 1. write-up
def build_writeup():
    S = styles(9.1)
    b = S["b"]
    doc = SimpleDocTemplate(str(DOCS / "Team9_Project35_Writeup.pdf"), pagesize=A4,
                            leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.3 * cm,
                            title=TITLE, author="Rohan Suresh, Satvik L")
    t1, bold = table1_rows()
    story = [
        Paragraph(TITLE, S["title"]),
        Paragraph(f"{COURSE}<br/>{TEAM}  |  {MEMBERS}", S["sub"]),

        Paragraph("1. Problem statement", S["h"]),
        Paragraph(
            "Maritime monitoring (illegal fishing, piracy, accident response) needs ships to be found "
            "in satellite imagery automatically. Following the Stanford CS229 project <i>Airbus Ship "
            "Detection – Traditional v.s. Convolutional Neural Network Approach</i> (Chen, Zheng, Zhou, "
            "2018), we treat it as <b>binary image classification</b>: given a satellite image, decide "
            "whether it contains at least one ship. The aim is to compare traditional machine-learning "
            "classifiers on hand-engineered features against convolutional neural networks that learn "
            "their features from pixels. It is hard because ships are tiny relative to the image and "
            "scenes vary widely (open sea, docks, clouds, haze, wakes).", b),

        Paragraph("2. Dataset", S["h"]),
        Paragraph(
            "Source: the Kaggle <b>Airbus Ship Detection Challenge</b> – 192,555 labelled 768×768 RGB "
            "satellite images (31 GB) and a CSV of run-length-encoded ship masks; an image with no mask "
            "has no ship (22.1% of all images contain one). The reference report used a 15k-image "
            f"subset; with a 1 GB download budget we used <b>{N:,} images</b>: a ready-made 494 MB "
            "subset of the competition's training images plus 515 images sampled uniformly at random "
            "(Kaggle rate-limits per-image downloads). Every label comes from the official competition "
            f"CSV. Our subset is ship-enriched – <b>{pct(SHIP)} ship / {pct(1 - SHIP)} no ship</b> – so "
            f"the trivial 'always predict ship' baseline is {pct(BASE)}, and we report precision, recall "
            "and specificity next to accuracy. Images are resized to 256×256×3 and split with "
            f"stratification in the report's 7:3:5 ratio: <b>{N_TR:,} train / {N_DV:,} dev / "
            f"{N_TE:,} test</b>. The test set is used only for the final numbers.", b),

        Paragraph("3. Approach", S["h"]),
        Paragraph(
            "<b>Feature engineering (traditional models).</b> Three global descriptors per image: "
            "<b>His</b> – an 8×8×8 colour histogram in HSV space (512 values, colour); <b>Ha</b> – 13 "
            "Haralick statistics of the grey-level co-occurrence matrix (texture); <b>Hu</b> – the 7 Hu "
            "moments (shape). As the 'no feature engineering' baseline the image is shrunk to 32×32×3 "
            "and flattened (3,072 raw pixel values). Features are standardised using training-set "
            "statistics.", b),
        Paragraph(
            "<b>Traditional classifiers.</b> Linear Discriminant Analysis, k-Nearest Neighbours (k = 5, "
            "Euclidean), Gaussian Naive Bayes, Random Forest (70 trees) and an RBF-kernel SVM, each "
            "trained on all 8 feature combinations (40 experiments), plus stratified 10-fold "
            "cross-validation on train + dev to measure stability.", b),
        Paragraph(
            f"<b>SIM-CNN (from scratch).</b> 4 × [3×3 convolution → ReLU → 2×2 max-pool] with "
            f"16/32/64/128 filters, then FC(128) → dropout → FC(1) → sigmoid; {P_SIM:,} trainable "
            "parameters, identical to the report. Trained with augmentation (flips, rotation, shift, "
            "zoom, shear, brightness). "
            f"<b>TL-CNN (transfer learning).</b> ImageNet-pretrained DenseNet-169 with all weights "
            f"frozen → global max-pool → BatchNorm → dropout → FC(128) → FC(1) → sigmoid; only the "
            f"{P_TL:,}-parameter head is trained (also identical to the report).", b),
        Paragraph(
            "<b>Training and evaluation.</b> Binary cross-entropy loss, Adam (lr 10<sup>-3</sup>, "
            "halved when the dev loss stalls), batch size 64, 30 epochs, keeping the checkpoint with "
            "the lowest dev loss. Because the classes are imbalanced, the decision threshold is chosen "
            "by scanning 0.1–0.9 on the dev set instead of assuming 0.5.", b),

        Paragraph("4. Implementation overview", S["h"]),
        Paragraph(
            "Python 3.13; scikit-learn for the traditional models, OpenCV and mahotas for features, "
            "PyTorch/torchvision for the CNNs; trained on a laptop CPU. The code is a small package "
            "(<font face='Courier'>src/</font>) with one module per stage: "
            "<font face='Courier'>download_data</font> (subset download, resize, labels, split) → "
            "<font face='Courier'>features</font> → <font face='Courier'>train_traditional</font> → "
            "<font face='Courier'>train_cnn</font> → <font face='Courier'>evaluate_cnn</font> → "
            "<font face='Courier'>make_figures</font>, plus <font face='Courier'>demo</font> for live "
            "prediction. <font face='Courier'>run_all.py</font> runs every stage; all settings live in "
            "<font face='Courier'>config.py</font>; random seeds are fixed and the exact image list and "
            "split are committed (<font face='Courier'>data/labels.csv</font>). Since the DenseNet "
            "backbone is frozen, its output for each image is computed once and cached, so the TL-CNN "
            "head trains in seconds.", b),

        Paragraph("5. Results", S["h"]),
        Paragraph(f"<b>Table 1.</b> Test accuracy of the traditional models for each feature set "
                  f"({N_TE:,} test images; best value in bold).", S["cap"]),
        grid(t1, col_widths=[1.3 * cm] + [1.85 * cm] * 8, font=7.8, bold_cells=bold),
        Spacer(1, 5),
        KeepTogether([
            Paragraph("<b>Table 2.</b> All models on the same test set (traditional models use all "
                      "three descriptors).", S["cap"]),
            grid(summary_rows(), col_widths=[4.1 * cm, 1.8 * cm, 1.8 * cm, 1.8 * cm, 6.2 * cm],
                 font=7.8)]),
        Spacer(1, 4),
        Table([[Image(str(FIG / "model_comparison.png"), width=8.7 * cm, height=4.72 * cm),
                Image(str(FIG / "confusion_matrices.png"), width=8.7 * cm, height=3.89 * cm)]],
              hAlign="CENTER"),
        Paragraph("<b>Figure 1.</b> Left: test accuracy of all models against the majority-class "
                  "baseline. Right: normalised confusion matrices of the two CNNs.", S["cap"]),

        Paragraph("6. Conclusions", S["h"]),
        bullets([
            f"<b>Random Forest is the best traditional model</b> ({pct(acc_all['RF'])} test accuracy, "
            f"10-fold CV {pct(cv['RF']['mean'])} ± {pct(cv['RF']['std'])}), in agreement with the "
            "reference report (93%).",
            f"<b>Feature engineering matters:</b> going from raw pixels to the three descriptors "
            f"changed accuracy by {', '.join(f'{a} {100 * FE_GAIN[a]:+.1f}' for a in ALGS)} points. "
            + (f"It helped {', '.join(fe_helped)}" if fe_helped else "")
            + (f" but hurt {', '.join(fe_hurt)}" if fe_hurt else "")
            + ". More features are not always better – the 512 correlated histogram bins break Naive "
              "Bayes's independence assumption, as the report also observed.",
            f"<b>CNNs:</b> {CNN_SENTENCE}; the best model overall was {best_overall[0]} "
            f"({pct(best_overall[1])}). "
            + ("Unlike the report, transfer learning beat the from-scratch network. A likely reason "
               f"is data size: we trained on {N_TR:,} images instead of 7,000, and a 4.3 M-parameter "
               "network trained from scratch needs far more data than a small head on features "
               "pre-learned from 1.2 M ImageNet images. SIM-CNN also scored below Random Forest."
               if tl["accuracy"] > sim["accuracy"] else
               "As in the report, the network trained from scratch on this specific imagery beat the "
               "frozen ImageNet features."),
            "<b>Limitations:</b> a small, ship-enriched subset (so absolute numbers are not directly "
            "comparable with the report's 84%-no-ship data), CPU-only training, a single run per "
            "model, SIM-CNN's dev loss was still falling at epoch 30 (it had not converged), and "
            "image-level classification only. <b>Future work:</b> train on the full "
            "dataset, fine-tune the backbone, and segment/localise the ships.",
        ], b),
    ]
    doc.build(story)


# ----------------------------------------------------------------- 2. slides
# Deliberately plain: black text on a white background, no colours, no figures.
BLACK = RGBColor(0, 0, 0)


def build_slides():
    from pptx.oxml.ns import qn
    from lxml import etree

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, lines, size=22, bold=False):
        tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(size), bold, BLACK, "Calibri"
            p.space_after = Pt(size * 0.5)

    def slide(title, bullets=None, size=22):
        s = prs.slides.add_slide(blank)
        text(s, 0.7, 0.45, 12, 0.9, title, size=32, bold=True)
        if bullets:
            text(s, 0.7, 1.5, 11.9, 5.6, ["•  " + b for b in bullets], size=size)
        return s

    def border(cell):
        tcPr = cell._tc.get_or_add_tcPr()
        for i, edge in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):   # must precede the fill
            ln = etree.Element(qn(edge), w="9525")
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr"), val="000000")
            tcPr.insert(i, ln)

    def table(s, rows, y=1.5, col_w=None, size=16):
        n_rows, n_cols = len(rows), len(rows[0])
        col_w = col_w or [11.9 / n_cols] * n_cols
        shape = s.shapes.add_table(n_rows, n_cols, Inches(0.7), Inches(y), Inches(sum(col_w)),
                                   Inches(0.45 * n_rows))
        t = shape.table
        t.first_row = t.horz_banding = False
        for c, w in enumerate(col_w):
            t.columns[c].width = Inches(w)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.text = str(val)
                f = cell.text_frame.paragraphs[0].font
                f.size, f.color.rgb, f.name, f.bold = Pt(size), BLACK, "Calibri", r == 0
                border(cell)

    # 1 - title
    s = prs.slides.add_slide(blank)
    text(s, 0.7, 2.2, 12, 1.2, "Airbus Ship Detection", size=44, bold=True)
    text(s, 0.7, 3.3, 12, 0.7, "Traditional Machine Learning vs. Convolutional Neural Networks", size=24)
    text(s, 0.7, 4.6, 12, 1.8, [COURSE, TEAM, MEMBERS], size=18)

    slide("Problem statement", [
        "Given a satellite image, decide whether it contains at least one ship (binary classification)",
        "Compare classical models on hand-engineered features with CNNs that learn features from pixels",
        "Reproduces a Stanford CS229 project (Chen, Zheng, Zhou, 2018)",
    ])

    slide("Dataset", [
        "Kaggle Airbus Ship Detection Challenge: 192,555 satellite images (768×768), 31 GB",
        "Labels from the official mask CSV: at least one ship mask = Ship, otherwise No ship",
        f"Our subset: {N:,} images, resized to 256×256",
        f"{pct(SHIP)} Ship / {pct(1 - SHIP)} No ship",
        f"Stratified 7:3:5 split: {N_TR:,} train / {N_DV:,} dev / {N_TE:,} test",
    ])

    slide("Approach", [
        "Hand-engineered features: colour histogram (512), Haralick texture (13), Hu moments (7)",
        "Classical models: LDA, k-NN (k = 5), Naive Bayes, Random Forest (70 trees), SVM (RBF)",
        f"SIM-CNN: 4 conv + max-pool layers, FC(128), FC(1); {P_SIM:,} parameters, trained from scratch",
        f"TL-CNN: frozen ImageNet DenseNet-169 + small classifier head; {P_TL:,} trained parameters",
        "CNN training: binary cross-entropy, Adam, batch 64, 30 epochs, best dev-loss checkpoint",
        "Decision threshold chosen on the dev set; final scores on the untouched test set",
    ], size=21)

    rows, _ = table1_rows()
    s = slide("Results – classical models (test accuracy)")
    table(s, rows, col_w=[1.3] + [1.325] * 8, size=15)

    def spec(a):
        d = trad["detail"][f"{a}|{ALL}"]
        n_pos = int(round(trad["ship_rate_test"] * N_TE))
        tp = d["recall"] * n_pos
        fp = tp / d["precision"] - tp if d["precision"] else 0
        return 1 - fp / (N_TE - n_pos)

    res = [["Model", "Accuracy", "Precision", "Recall", "Specificity", "F1"]]
    for a, name in (("LDA", "LDA"), ("KNN", "k-NN"), ("NB", "Naive Bayes"),
                    ("RF", "Random Forest"), ("SVM", "SVM")):
        d = trad["detail"][f"{a}|{ALL}"]
        res.append([name, pct(d["accuracy"]), pct(d["precision"]), pct(d["recall"]),
                    pct(spec(a)), pct(d["f1"])])
    for m, name in ((tl, "TL-CNN"), (sim, "SIM-CNN")):
        f1 = 2 * m["precision"] * m["recall"] / (m["precision"] + m["recall"])
        res.append([name, pct(m["accuracy"]), pct(m["precision"]), pct(m["recall"]),
                    pct(m["specificity"]), pct(f1)])
    res.append(["Always 'Ship' (baseline)", pct(BASE), pct(BASE), "100.0%", "0.0%",
                pct(2 * BASE / (1 + BASE))])
    s = slide("Results – all models (test set)")
    table(s, res, col_w=[3.4, 1.7, 1.7, 1.7, 1.7, 1.7], size=16)

    s = slide("Comparison with the Stanford report")
    table(s, [["", "Stanford report", "Our project"],
              ["Images (train / dev / test)", "7,000 / 3,000 / 5,000", f"{N_TR:,} / {N_DV:,} / {N_TE:,}"],
              ["Best classical model", "Random Forest, 93%", f"Random Forest, {pct(acc_all['RF'])}"],
              ["TL-CNN accuracy", "90%", pct(tl["accuracy"])],
              ["SIM-CNN accuracy", "94%", pct(sim["accuracy"])]],
          col_w=[4.3, 3.8, 3.8], size=18)

    prs.save(str(DOCS / "Team9_Project35_Slides.pptx"))


if __name__ == "__main__":
    from docs.viva_guide import build_viva_guide
    build_writeup()
    build_slides()
    build_viva_guide(globals())
    print("built:", *[p.name for p in sorted(DOCS.glob("Team9_*"))], sep="\n  ")
